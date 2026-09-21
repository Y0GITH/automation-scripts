import os
import re
import time
import pyodbc
from concurrent.futures import ThreadPoolExecutor, as_completed
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ============================================================
# CONFIGURATION
# ============================================================

# Default number of rows to capture per result set.
# The user is asked for this at runtime; if they leave it
# blank (or enter something invalid), this default is used.
DEFAULT_MAX_ROWS = 10

# Number of SQL files that can run at the same time.
# Be careful if your SQL files modify the same tables.
MAX_WORKERS = 5

# Change this if your machine has a different driver.
ODBC_DRIVER = "ODBC Driver 18 for SQL Server"


# ============================================================
# GENERAL FUNCTIONS
# ============================================================


def clean_path(path):
    """Remove quotes and whitespace from a path."""
    return path.strip().strip('"').strip("'")


def find_sql_files(root_folder):
    """Recursively find every .sql file."""
    sql_files = []

    for root, dirs, files in os.walk(root_folder):
        for file in files:
            if file.lower().endswith(".sql"):
                full_path = os.path.join(root, file)
                sql_files.append(full_path)

    sql_files.sort()

    return sql_files


def split_sql_batches(sql_text):
    """
    Split SQL script using SSMS-style GO statements.

    GO must appear on a line by itself, optionally followed
    by a number.
    """

    pattern = r"(?im)^\s*GO\s*(?:\d+)?\s*(?:--.*)?$"

    batches = re.split(pattern, sql_text)

    batches = [batch.strip() for batch in batches if batch.strip()]

    return batches


def format_value(value):
    """Convert database values into readable text."""

    if value is None:
        return "NULL"

    if isinstance(value, bytes):
        try:
            return value.decode("utf-8")
        except Exception:
            return str(value)

    return str(value)


def ask_max_rows():
    """
    Ask the user how many rows to capture per result set
    (like selecting TOP N). If nothing is entered, or the
    input isn't a valid positive number, fall back to the
    default.
    """

    raw = input(
        f"\nHow many rows would you like to capture per result set?\n"
        f"(Press Enter for default: {DEFAULT_MAX_ROWS})\n> "
    ).strip()

    if raw == "":
        return DEFAULT_MAX_ROWS

    try:
        value = int(raw)

        if value <= 0:
            print(f"Value must be positive. Using default: {DEFAULT_MAX_ROWS}")
            return DEFAULT_MAX_ROWS

        return value

    except ValueError:
        print(f"Invalid number. Using default: {DEFAULT_MAX_ROWS}")
        return DEFAULT_MAX_ROWS


# ============================================================
# SQL CONNECTION
# ============================================================


def create_connection(server, database, auth_type, username=None, password=None):

    if auth_type == "windows":

        connection_string = (
            f"DRIVER={{{ODBC_DRIVER}}};"
            f"SERVER={server};"
            f"DATABASE={database};"
            f"Trusted_Connection=yes;"
            f"TrustServerCertificate=yes;"
        )

    else:

        connection_string = (
            f"DRIVER={{{ODBC_DRIVER}}};"
            f"SERVER={server};"
            f"DATABASE={database};"
            f"UID={username};"
            f"PWD={password};"
            f"TrustServerCertificate=yes;"
        )

    return pyodbc.connect(connection_string, timeout=30)


# ============================================================
# EXECUTE ONE SQL FILE
# ============================================================


def execute_sql_file(
    file_path,
    root_folder,
    server,
    database,
    auth_type,
    username=None,
    password=None,
    max_rows=DEFAULT_MAX_ROWS,
):

    start_time = time.time()

    relative_path = os.path.relpath(file_path, root_folder)

    folder_path = os.path.dirname(relative_path)

    if folder_path == "":
        folder_path = "(Root Folder)"

    file_name = os.path.basename(file_path)

    result = {
        "file_path": file_path,
        "relative_path": relative_path,
        "folder": folder_path,
        "file_name": file_name,
        "status": "FAILED",
        "execution_time": 0,
        "batches": [],
        "error": None,
    }

    connection = None
    cursor = None

    try:

        # ----------------------------------------------------
        # Read SQL file
        # ----------------------------------------------------

        with open(file_path, "r", encoding="utf-8-sig", errors="replace") as f:

            sql_text = f.read()

        if not sql_text.strip():

            result["status"] = "EMPTY"

            return result

        # ----------------------------------------------------
        # Split GO batches
        # ----------------------------------------------------

        batches = split_sql_batches(sql_text)

        if not batches:

            result["status"] = "EMPTY"

            return result

        # ----------------------------------------------------
        # Create a separate SQL connection for this file
        # ----------------------------------------------------

        connection = create_connection(server, database, auth_type, username, password)

        cursor = connection.cursor()

        # ----------------------------------------------------
        # Execute every batch
        # ----------------------------------------------------

        for batch_number, batch in enumerate(batches, start=1):

            batch_result = {
                "batch_number": batch_number,
                "code": batch,
                "result_sets": [],
                "messages": [],
                "error": None,
            }

            try:

                cursor.execute(batch)

                result_set_number = 0

                # ------------------------------------------------
                # Capture ALL result sets
                # ------------------------------------------------

                while True:

                    if cursor.description:

                        result_set_number += 1

                        columns = [column[0] for column in cursor.description]

                        rows = cursor.fetchmany(max_rows)

                        formatted_rows = []

                        for row in rows:

                            formatted_rows.append(
                                [format_value(value) for value in row]
                            )

                        # Try to determine total rows.
                        #
                        # fetchmany only retrieves max_rows, so
                        # rows_captured below is the number captured,
                        # not necessarily total rows returned by SQL
                        # Server.
                        #
                        # We therefore mark it clearly.

                        result_set = {
                            "result_set_number": result_set_number,
                            "columns": columns,
                            "rows": formatted_rows,
                            "rows_captured": len(formatted_rows),
                            "truncated": len(formatted_rows) == max_rows,
                        }

                        batch_result["result_sets"].append(result_set)

                    # ------------------------------------------------
                    # Move to next result set
                    # ------------------------------------------------

                    try:

                        has_next = cursor.nextset()

                    except pyodbc.Error:

                        has_next = False

                    if not has_next:
                        break

                # ------------------------------------------------
                # Commit if SQL changed data
                # ------------------------------------------------

                connection.commit()

            except Exception as batch_error:

                batch_result["error"] = str(batch_error)

                # Rollback this connection if possible
                try:
                    connection.rollback()
                except Exception:
                    pass

            result["batches"].append(batch_result)

        # ----------------------------------------------------
        # Determine overall status
        # ----------------------------------------------------

        errors = [batch for batch in result["batches"] if batch["error"]]

        if errors:

            result["status"] = "COMPLETED WITH ERRORS"

            result["error"] = "\n".join(
                f"Batch {b['batch_number']}: {b['error']}" for b in errors
            )

        else:

            result["status"] = "SUCCESS"

    except Exception as e:

        result["status"] = "FAILED"
        result["error"] = str(e)

    finally:

        if cursor:

            try:
                cursor.close()
            except Exception:
                pass

        if connection:

            try:
                connection.close()
            except Exception:
                pass

        result["execution_time"] = time.time() - start_time

    return result


# ============================================================
# DOCUMENT FORMATTING FUNCTIONS
# ============================================================


def set_cell_text(cell, text, bold=False):

    cell.text = ""

    paragraph = cell.paragraphs[0]

    run = paragraph.add_run(str(text))

    run.bold = bold
    run.font.size = Pt(9)

    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def shade_cell(cell, fill="D9EAF7"):

    tc_pr = cell._tc.get_or_add_tcPr()

    shd = OxmlElement("w:shd")

    shd.set(qn("w:fill"), fill)

    tc_pr.append(shd)


def set_repeat_table_header(row):

    tr_pr = row._tr.get_or_add_trPr()

    tbl_header = OxmlElement("w:tblHeader")

    tbl_header.set(qn("w:val"), "true")

    tr_pr.append(tbl_header)


def add_code_block(document, code):

    paragraph = document.add_paragraph()

    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.25)

    run = paragraph.add_run(code)

    run.font.name = "Consolas"
    run.font.size = Pt(8)

    return paragraph


def add_result_table(document, columns, rows):

    if not columns:

        document.add_paragraph("No columns returned.")

        return

    table = document.add_table(rows=1, cols=len(columns))

    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    table.style = "Table Grid"

    header = table.rows[0]

    set_repeat_table_header(header)

    for i, column in enumerate(columns):

        set_cell_text(header.cells[i], column, bold=True)

        shade_cell(header.cells[i])

    # --------------------------------------------------------
    # Add rows
    # --------------------------------------------------------

    for row in rows:

        cells = table.add_row().cells

        for i, value in enumerate(row):

            set_cell_text(cells[i], value)

    document.add_paragraph()


# ============================================================
# CREATE WORD DOCUMENT
# ============================================================


def create_document(results, root_folder, output_path, max_rows=DEFAULT_MAX_ROWS):

    document = Document()

    # --------------------------------------------------------
    # Page setup
    # --------------------------------------------------------

    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    # --------------------------------------------------------
    # Default font
    # --------------------------------------------------------

    styles = document.styles

    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    # --------------------------------------------------------
    # Title
    # --------------------------------------------------------

    title = document.add_heading("SQL Execution Documentation", level=0)

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph(f"Root Folder: {root_folder}")

    document.add_paragraph(f"Total SQL Files: {len(results)}")

    document.add_paragraph(f"Maximum Rows Captured Per Result Set: {max_rows}")

    document.add_page_break()

    # --------------------------------------------------------
    # Sort results by folder and file
    # --------------------------------------------------------

    results.sort(key=lambda x: (x["folder"].lower(), x["file_name"].lower()))

    # --------------------------------------------------------
    # Track folders
    # --------------------------------------------------------

    current_folder = None

    for result in results:

        folder = result["folder"]

        # ----------------------------------------------------
        # New folder
        # ----------------------------------------------------

        if folder != current_folder:

            if current_folder is not None:

                document.add_page_break()

            current_folder = folder

            document.add_heading(f"Folder: {folder}", level=1)

        # ----------------------------------------------------
        # File name
        # ----------------------------------------------------

        document.add_heading(f"File: {result['file_name']}", level=2)

        # ----------------------------------------------------
        # File path
        # ----------------------------------------------------

        document.add_paragraph(f"Path: {result['relative_path']}")

        # ----------------------------------------------------
        # Status
        # ----------------------------------------------------

        status_paragraph = document.add_paragraph()

        status_run = status_paragraph.add_run(f"Execution Status: {result['status']}")

        status_run.bold = True

        document.add_paragraph(
            f"Execution Time: " f"{result['execution_time']:.2f} seconds"
        )

        # ----------------------------------------------------
        # File-level error
        # ----------------------------------------------------

        if result["error"]:

            document.add_heading("Error", level=3)

            error_paragraph = document.add_paragraph()

            error_run = error_paragraph.add_run(result["error"])

            error_run.font.name = "Consolas"
            error_run.font.size = Pt(8)

        # ----------------------------------------------------
        # Batches
        # ----------------------------------------------------

        for batch in result["batches"]:

            document.add_heading(f"SQL Batch {batch['batch_number']}", level=3)

            # ------------------------------------------------
            # SQL CODE
            # ------------------------------------------------

            document.add_heading("SQL Code", level=4)

            add_code_block(document, batch["code"])

            # ------------------------------------------------
            # Batch error
            # ------------------------------------------------

            if batch["error"]:

                document.add_heading("Batch Error", level=4)

                error_paragraph = document.add_paragraph()

                error_run = error_paragraph.add_run(batch["error"])

                error_run.font.name = "Consolas"
                error_run.font.size = Pt(8)

            # ------------------------------------------------
            # Result sets
            # ------------------------------------------------

            if batch["result_sets"]:

                for result_set in batch["result_sets"]:

                    document.add_heading(
                        f"Result Set " f"{result_set['result_set_number']}", level=4
                    )

                    document.add_paragraph(
                        f"Rows captured: " f"{result_set['rows_captured']}"
                    )

                    if result_set["truncated"]:

                        document.add_paragraph(
                            f"Only the first {max_rows} rows "
                            f"are included in this document."
                        )

                    add_result_table(
                        document, result_set["columns"], result_set["rows"]
                    )

            else:

                document.add_paragraph("No tabular result set returned.")

        document.add_paragraph("-" * 80)

    # --------------------------------------------------------
    # Summary
    # --------------------------------------------------------

    document.add_page_break()

    document.add_heading("Execution Summary", level=1)

    success = sum(1 for r in results if r["status"] == "SUCCESS")

    errors = sum(
        1 for r in results if r["status"] in ["FAILED", "COMPLETED WITH ERRORS"]
    )

    empty = sum(1 for r in results if r["status"] == "EMPTY")

    document.add_paragraph(f"Total Files: {len(results)}")

    document.add_paragraph(f"Successful: {success}")

    document.add_paragraph(f"Failed / Errors: {errors}")

    document.add_paragraph(f"Empty Files: {empty}")

    document.save(output_path)


# ============================================================
# MAIN PROGRAM
# ============================================================


def main():

    print("=" * 70)
    print("       SQL PROJECT EXECUTION & DOCUMENTATION TOOL")
    print("=" * 70)

    print()

    # --------------------------------------------------------
    # Ask for root folder
    # --------------------------------------------------------

    root_folder = input("Enter the root folder containing your SQL files:\n> ")

    root_folder = clean_path(root_folder)

    if not os.path.isdir(root_folder):

        print()
        print("ERROR: Folder does not exist.")

        return

    # --------------------------------------------------------
    # Find SQL files
    # --------------------------------------------------------

    print()
    print("Searching for SQL files...")

    sql_files = find_sql_files(root_folder)

    if not sql_files:

        print()
        print("No .sql files were found.")

        return

    print()
    print(f"Found {len(sql_files)} SQL files.")

    print()

    for i, file in enumerate(sql_files, start=1):

        relative = os.path.relpath(file, root_folder)

        print(f"{i:03d}. {relative}")

    # --------------------------------------------------------
    # ROW LIMIT (like selecting TOP N)
    # --------------------------------------------------------

    max_rows = ask_max_rows()

    print(f"\nRows per result set set to: {max_rows}")

    # --------------------------------------------------------
    # SQL SERVER CONNECTION DETAILS
    # --------------------------------------------------------

    print()
    print("=" * 70)
    print("SQL SERVER CONNECTION")
    print("=" * 70)

    server = input("\nEnter SQL Server name:\n> ").strip()

    database = input("\nEnter database name:\n> ").strip()

    print()
    print("Authentication:")
    print("1 = Windows Authentication")
    print("2 = SQL Server Authentication")

    auth_choice = input("> ").strip()

    if auth_choice == "2":

        auth_type = "sql"

        username = input("\nSQL Username:\n> ").strip()

        password = input("\nSQL Password:\n> ").strip()

    else:

        auth_type = "windows"

        username = None
        password = None

    # --------------------------------------------------------
    # Confirm before starting
    # --------------------------------------------------------

    print()
    print("=" * 70)
    print("EXECUTION PLAN")
    print("=" * 70)

    print(f"SQL files       : {len(sql_files)}")

    print(f"Parallel workers: {MAX_WORKERS}")

    print(f"Rows/result set : {max_rows}")

    print("\nAll files will now be executed.")

    print("The program will wait until ALL files finish.")

    print()

    confirmation = input("Start execution? (Y/N): ").strip().lower()

    if confirmation != "y":

        print("Execution cancelled.")

        return

    # --------------------------------------------------------
    # Execute files in parallel
    # --------------------------------------------------------

    print()
    print("=" * 70)
    print("STARTING SQL EXECUTION")
    print("=" * 70)

    all_results = []

    overall_start = time.time()

    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:

        future_to_file = {}

        for file_path in sql_files:

            future = executor.submit(
                execute_sql_file,
                file_path,
                root_folder,
                server,
                database,
                auth_type,
                username,
                password,
                max_rows,
            )

            future_to_file[future] = file_path

        completed = 0

        for future in as_completed(future_to_file):

            file_path = future_to_file[future]

            completed += 1

            try:

                result = future.result()

                all_results.append(result)

                print(
                    f"[{completed}/{len(sql_files)}] "
                    f"{result['relative_path']} "
                    f"--> "
                    f"{result['status']} "
                    f"({result['execution_time']:.2f}s)"
                )

            except Exception as e:

                print(
                    f"[{completed}/{len(sql_files)}] "
                    f"{file_path} "
                    f"--> UNEXPECTED ERROR: {e}"
                )

    overall_time = time.time() - overall_start

    # --------------------------------------------------------
    # ALL EXECUTIONS FINISHED
    # --------------------------------------------------------

    print()
    print("=" * 70)
    print("ALL SQL FILES HAVE FINISHED")
    print("=" * 70)

    print(f"Total execution time: " f"{overall_time / 60:.2f} minutes")

    # --------------------------------------------------------
    # Ask where to save document
    # --------------------------------------------------------

    print()
    print("=" * 70)
    print("SAVE DOCUMENT")
    print("=" * 70)

    output_path = input(
        "\nEnter complete path for the Word document:\n"
        "Example: C:\\Users\\Lenovo\\Documents\\SQL_Report.docx\n> "
    )

    output_path = clean_path(output_path)

    # Automatically add .docx
    if not output_path.lower().endswith(".docx"):

        output_path += ".docx"

    # --------------------------------------------------------
    # Create output folder if necessary
    # --------------------------------------------------------

    output_directory = os.path.dirname(output_path)

    if output_directory:

        os.makedirs(output_directory, exist_ok=True)

    # --------------------------------------------------------
    # Create document
    # --------------------------------------------------------

    print()
    print("Creating Word documentation...")

    create_document(all_results, root_folder, output_path, max_rows)

    # --------------------------------------------------------
    # Finished
    # --------------------------------------------------------

    print()
    print("=" * 70)
    print("DOCUMENTATION COMPLETE")
    print("=" * 70)

    print()
    print(f"Document saved to:\n{output_path}")

    print()


if __name__ == "__main__":
    main()
