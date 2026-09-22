"""
doc_to_speech.py

Reads a document (.docx, .txt, or .pdf), intelligently narrates its content
(including tables, read row-by-row with structure announced first), normalizes
numbers/currency/percentages into natural spoken English, and converts the
result into a compact MP3 using Microsoft Edge's neural TTS (edge-tts).

Requirements:
    pip install python-docx edge-tts num2words pypdf

Usage:
    python doc_to_speech.py
    (the script will interactively ask for the input file, voice, and output location)
"""

import asyncio
import os
import re
import sys

# ---------- Dependency check ----------
try:
    import docx  # python-docx
except ImportError:
    docx = None

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    import edge_tts
except ImportError:
    print("Missing dependency 'edge-tts'. Install with: pip install edge-tts")
    sys.exit(1)

try:
    from num2words import num2words
except ImportError:
    print("Missing dependency 'num2words'. Install with: pip install num2words")
    sys.exit(1)


# ---------- Voice options ----------
VOICES = {
    "1": {"label": "Male", "id": "en-US-GuyNeural"},
    "2": {"label": "Female", "id": "en-US-JennyNeural"},
}


# ---------- Number / currency / percentage normalization ----------

CURRENCY_RE = re.compile(r"\$\s?(\d[\d,]*)(\.(\d{1,2}))?")
PERCENT_RE = re.compile(r"(\d+(\.\d+)?)\s?%")
PLAIN_NUMBER_RE = re.compile(r"(?<![\w.])(\d[\d,]*)(\.(\d+))?(?![\w])")


def _int_to_words(n_str: str) -> str:
    n = int(n_str.replace(",", ""))
    return num2words(n)


def normalize_currency(match: re.Match) -> str:
    whole = match.group(1)
    frac = match.group(3)
    words = (
        _int_to_words(whole)
        + " dollar"
        + ("" if whole.replace(",", "") == "1" else "s")
    )
    if frac:
        cents = int(frac.ljust(2, "0"))
        if cents > 0:
            words += " and " + num2words(cents) + " cent" + ("" if cents == 1 else "s")
    return words


def normalize_percent(match: re.Match) -> str:
    value = match.group(1)
    if "." in value:
        whole, dec = value.split(".")
        spoken = (
            num2words(int(whole)) + " point " + " ".join(num2words(int(d)) for d in dec)
        )
    else:
        spoken = num2words(int(value))
    return spoken + " percent"


def normalize_plain_number(match: re.Match) -> str:
    whole = match.group(1)
    frac = match.group(3)
    try:
        if frac:
            whole_words = num2words(int(whole.replace(",", "")))
            dec_words = " ".join(num2words(int(d)) for d in frac)
            return f"{whole_words} point {dec_words}"
        else:
            return num2words(int(whole.replace(",", "")))
    except ValueError:
        return match.group(0)


def normalize_text(text: str) -> str:
    """Convert currency, percentages, and plain numbers into natural spoken words."""
    text = CURRENCY_RE.sub(normalize_currency, text)
    text = PERCENT_RE.sub(normalize_percent, text)
    text = PLAIN_NUMBER_RE.sub(normalize_plain_number, text)
    return text


# ---------- Document readers ----------


def read_docx(path: str) -> str:
    """Walk a .docx body in order, narrating paragraphs and tables distinctly."""
    document = docx.Document(path)
    spoken_parts = []

    # Walk the document body in order (paragraphs and tables interleaved)
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]

        if tag == "p":
            para = next((p for p in document.paragraphs if p._p == child), None)
            if para is None or not para.text.strip():
                continue
            style_name = para.style.name if para.style is not None else ""
            style = (style_name or "").lower()
            text = para.text.strip()
            if "heading" in style or "title" in style:
                spoken_parts.append(f"Heading: {text}.")
            elif "list" in style:
                spoken_parts.append(f"Bullet point: {text}.")
            else:
                spoken_parts.append(text)

        elif tag == "tbl":
            table = next((t for t in document.tables if t._tbl == child), None)
            if table is None:
                continue
            spoken_parts.append(describe_table(table))

    return "\n\n".join(spoken_parts)


def describe_table(table) -> str:
    """Narrate a docx table: structure summary, then row-by-row with column names."""
    rows = table.rows
    if not rows:
        return ""

    headers = [cell.text.strip() for cell in rows[0].cells]
    n_cols = len(headers)
    n_rows = len(rows) - 1

    lines = [
        f"Table with {n_cols} columns and {n_rows} data rows. "
        f"The columns are: {', '.join(h if h else f'column {i+1}' for i, h in enumerate(headers))}."
    ]

    for idx, row in enumerate(rows[1:], start=1):
        cells = [cell.text.strip() for cell in row.cells]
        row_desc = f"Row {idx}: " + "; ".join(
            f"{headers[i] if headers[i] else f'column {i+1}'} is {cells[i] if cells[i] else 'blank'}"
            for i in range(min(len(headers), len(cells)))
        )
        lines.append(row_desc + ".")

    return " ".join(lines)


def read_txt(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def read_pdf(path: str) -> str:
    if PdfReader is None:
        print("Missing dependency 'pypdf'. Install with: pip install pypdf")
        sys.exit(1)
    reader = PdfReader(path)
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def load_document(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".docx":
        if docx is None:
            print(
                "Missing dependency 'python-docx'. Install with: pip install python-docx"
            )
            sys.exit(1)
        return read_docx(path)
    elif ext == ".txt":
        return read_txt(path)
    elif ext == ".pdf":
        return read_pdf(path)
    else:
        print(f"Unsupported file type: {ext}. Supported: .docx, .txt, .pdf")
        sys.exit(1)


# ---------- TTS ----------


async def synthesize(text: str, voice_id: str, output_path: str):
    communicate = edge_tts.Communicate(text, voice_id, rate="+0%")
    await communicate.save(output_path)


# ---------- Interactive flow ----------


def ask_input_path() -> str:
    path = (
        input("Enter the path to the document (.docx, .txt, .pdf): ").strip().strip('"')
    )
    while not os.path.isfile(path):
        print("File not found.")
        path = input("Enter a valid file path: ").strip().strip('"')
    return path


def ask_voice() -> str:
    print("\nSelect voice:")
    for key, v in VOICES.items():
        print(f"{key}. {v['label']}")
    choice = input("Enter 1 or 2: ").strip()
    while choice not in VOICES:
        choice = input("Please enter 1 (Male) or 2 (Female): ").strip()
    return VOICES[choice]["id"]


def ask_output_path(default_dir: str, default_name: str) -> str:
    print(f"\nWhere should the audio file be saved?")
    print(f"(Press Enter to use default: {os.path.join(default_dir, default_name)})")
    save_dir = input("Folder path: ").strip().strip('"')
    if not save_dir:
        save_dir = default_dir
    os.makedirs(save_dir, exist_ok=True)

    filename = input(f"Filename (Enter for '{default_name}'): ").strip()
    if not filename:
        filename = default_name
    if not filename.lower().endswith(".mp3"):
        filename += ".mp3"

    return os.path.join(save_dir, filename)


def main():
    print("=== Document to Speech Converter ===\n")

    input_path = ask_input_path()
    print("\nReading document...")
    raw_text = load_document(input_path)

    if not raw_text.strip():
        print("No readable text found in this document.")
        sys.exit(1)

    print("Normalizing numbers, currency, and percentages for natural speech...")
    spoken_text = normalize_text(raw_text)

    voice_id = ask_voice()

    default_dir = os.path.dirname(os.path.abspath(input_path))
    default_name = os.path.splitext(os.path.basename(input_path))[0] + ".mp3"
    output_path = ask_output_path(default_dir, default_name)

    print(
        f"\nGenerating audio ({len(spoken_text)} characters of text)... this may take a moment."
    )
    asyncio.run(synthesize(spoken_text, voice_id, output_path))

    size_kb = os.path.getsize(output_path) / 1024
    print(f"\nDone. Audio saved to: {output_path} ({size_kb:.1f} KB)")


if __name__ == "__main__":
    main()
